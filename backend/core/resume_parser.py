import os
from docx import Document
from pypdf import PdfReader

def extract_resume_text(file_path):
    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".docx":
        return extract_docx_text(file_path)
    elif extension == ".pdf":
        return extract_pdf_text(file_path)
    else:
        raise ValueError(f"Unsupported file format: {extension}")
    
def extract_docx_text(file_path):
    doc = Document(file_path)

    text = []

    # Paragraphs
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)

    # Tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text.append(cell.text)

    return "\n".join(text)
    
def extract_pdf_text(file_path):
    reader = PdfReader(file_path)
    text = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text.append(page_text)

    return "\n".join(text)
